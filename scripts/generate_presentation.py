"""
Generate a professional, high-impact PowerPoint presentation for Hello Farmer (ሄሎ ፋርመር).
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path: str = "Hello_Farmer_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette Tokens
    C_PRIMARY_DARK = RGBColor(15, 56, 42)      # Deep Forest
    C_PRIMARY = RGBColor(30, 86, 66)           # Forest Green
    C_ACCENT_GOLD = RGBColor(234, 179, 8)      # Sun Gold
    C_ACCENT_GREEN = RGBColor(34, 197, 94)     # Leaf Green
    C_BG_LIGHT = RGBColor(248, 250, 249)       # Soft off-white
    C_CARD_BG = RGBColor(255, 255, 255)        # Pure White
    C_CARD_BORDER = RGBColor(226, 232, 240)    # Slate 200
    C_TEXT_MAIN = RGBColor(15, 23, 42)         # Slate 900
    C_TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500
    C_WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category_text="HELLO FARMER • ሄሎ ፋርመር"):
        # Category pill / banner
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = C_PRIMARY_DARK

    def add_card(slide, left, top, width, height, title, items, badge="", bg_color=C_CARD_BG, title_color=C_PRIMARY_DARK):
        # Card shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = C_CARD_BORDER
        shape.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        if badge:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge.upper()
            p_badge.font.size = Pt(9)
            p_badge.font.bold = True
            p_badge.font.color.rgb = C_PRIMARY
            p_badge.space_after = Pt(4)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.space_after = Pt(10)

        for it in items:
            p = tf.add_paragraph()
            p.text = f"• {it}" if not it.startswith("  ") else f"    {it}"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Cover)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_PRIMARY_DARK
    bg1.line.fill.background()

    # Title Container
    s1_title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.2))
    tf1 = s1_title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "NATIONAL AGRI-VOICE INTELLIGENCE PLATFORM • ETHIOPIA"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_ACCENT_GOLD
    p_badge.space_after = Pt(12)

    p_h1 = tf1.add_paragraph()
    p_h1.text = "Hello Farmer | ሄሎ ፋርመር"
    p_h1.font.size = Pt(44)
    p_h1.font.bold = True
    p_h1.font.color.rgb = C_WHITE
    p_h1.space_after = Pt(14)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Empowering Ethiopian Smallholder Farmers with Multilingual Voice AI\nBridging Literacy & Digital Barriers via National Hotline 8028 & Modern Web App"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(209, 231, 221)

    # Team & Metadata Footer
    s1_foot = s1.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.2))
    tff = s1_foot.text_frame
    pf1 = tff.paragraphs[0]
    pf1.text = "Project Contributors: Yeabsira Girma • KalMesKeb • Yohannes Girma"
    pf1.font.size = Pt(13)
    pf1.font.bold = True
    pf1.font.color.rgb = C_WHITE
    pf2 = tff.add_paragraph()
    pf2.text = "Languages Supported: Amharic (አማርኛ) • Afaan Oromo • Tigrinya (ትግርኛ) • Somali (Soomaali) • English"
    pf2.font.size = Pt(11)
    pf2.font.color.rgb = C_ACCENT_GOLD

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Rural Reality
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Core Problem: Rural Agricultural Information Gap")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "80%+ Agrarian Population",
        [
            "Agriculture forms the backbone of Ethiopia's economy and food security.",
            "Millions of smallholder farmers operate on rain-fed plots.",
            "Crucial need for timely fertilizer, pest, and planting decisions."
        ],
        badge="Context & Scale"
    )

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "The Digital & Literacy Divide",
        [
            "Low smartphone penetration in rural woredas.",
            "High rural illiteracy rates; text-based apps and SMS fail to engage.",
            "High linguistic diversity with complex morphology (Ge'ez and Latin scripts)."
        ],
        badge="Key Barrier"
    )

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Costly Agronomic Errors",
        [
            "Misdiagnosis of crop diseases (e.g., maize leaf blight, teff rust).",
            "Inappropriate pesticide usage and dangerous chemical dosage.",
            "Limited extension agents (DA ratio often 1 agent to 1,000+ farmers)."
        ],
        badge="Consequence"
    )

    # -------------------------------------------------------------
    # SLIDE 3: The Solution - Dual Channel Architecture
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "The Hello Farmer Solution: Inclusive Dual-Channel AI")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Channel 1: National Voice Hotline 8028",
        [
            "Zero Internet Needed: Runs over standard 2G/GSM feature phones.",
            "Toll-free shortcode 8028 integrated with Asterisk PBX.",
            "Natural spoken conversation: Farmer speaks in Amharic or Afaan Oromo.",
            "Real-time voice activity detection (VAD) and barge-in capability.",
            "Streams back natural synthesized voice advice in under 2 seconds."
        ],
        badge="Telephony (Feature Phone)"
    )

    add_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Channel 2: Modern Web Application",
        [
            "Voice-In / Voice-Out: Browser microphone speech capture.",
            "Audio File Upload: Test or upload recordings without mic hardware.",
            "Interactive Yield Calculator: Calibrated for teff, maize, wheat & barley.",
            "Hyperlocal Weather Risk Engine: Spraying window safety analyzer.",
            "Extension Dashboard: Call tracking, pest trend analytics, and SMS logs."
        ],
        badge="Smartphone & Desktop Web"
    )

    # -------------------------------------------------------------
    # SLIDE 4: End-to-End Pipeline & System Flow
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "End-to-End Technical Architecture")

    steps = [
        ("1. Speech Input", ["Telephone 8kHz AudioSocket or Web 16kHz audio stream", "Noise reduction & VAD"]),
        ("2. Speech Recognition", ["Google Gemini Audio STT or Groq Whisper Large-v3", "Transcribes to Ge'ez/Latin"]),
        ("3. Knowledge RAG", ["PostgreSQL pgvector database", "MoA ISFM manuals + BGE-M3 & Gemini Embeddings"]),
        ("4. LLM Reasoning", ["Groq Qwen-2.5-32B ultra-low latency engine", "Safety & pesticide guardrails"]),
        ("5. Neural Voice Output", ["Microsoft Edge Neural TTS (Mekdes / Ameha voices)", "AudioSocket 8kHz / Web WAV streaming"])
    ]

    card_w = Inches(2.25)
    gap = Inches(0.18)
    for idx, (title, items) in enumerate(steps):
        left_pos = Inches(0.8) + idx * (card_w + gap)
        add_card(s4, left_pos, Inches(1.8), card_w, Inches(5.0),
            title,
            items,
            badge=f"Step 0{idx+1}"
        )

    # -------------------------------------------------------------
    # SLIDE 5: The Working AI Engine Breakdown
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Multi-Model AI Stack: Proven & Operational")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Speech AI (STT & TTS)",
        [
            "Gemini Audio STT: Multimodal direct speech-to-text in Amharic & Oromo.",
            "Groq Whisper Large-v3: Sub-second cloud ASR.",
            "Local HuggingFace Whisper: Offline fallback for edge servers.",
            "Edge Neural TTS: Expressive Amharic neural speech (am-ET-MekdesNeural)."
        ],
        badge="Voice Intelligence"
    )

    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Reasoning & LLMs",
        [
            "Groq Qwen-2.5-32B: Fast, accurate agricultural problem-solving.",
            "Gemini 1.5 Flash / Pro: Long-context grounding & cross-lingual glosses.",
            "OpenRouter Nemotron: Cloud redundancy.",
            "Ollama (Llama3 / Qwen): Offline local inference capable."
        ],
        badge="Cognitive Core"
    )

    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Grounded Knowledge & RAG",
        [
            "Vector DB: PostgreSQL 16 with pgvector extension.",
            "Embeddings: Gemini text-embedding-004 & BAAI/bge-m3.",
            "Content: MoA Integrated Soil Fertility Management (ISFM) & pesticide guides.",
            "Live Web Search RAG: Real-time weather and commodity prices."
        ],
        badge="Factual Accuracy"
    )

    # -------------------------------------------------------------
    # SLIDE 6: Safety Guardrails & Agronomic Validation
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Safety Guardrails: Preventing Lethal Misinformation")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Pesticide & Chemical Safety Filter",
        [
            "Banned Substance Blocker: Strictly intercepts references to banned chemicals (e.g. DDT, Aldrin, Endosulfan).",
            "Dosage Validation: Checks chemical concentration and safety intervals (Pre-Harvest Interval - PHI).",
            "Personal Protective Equipment (PPE): Enforces protective handling warnings before spray recommendations.",
            "Medical Isolation: Redirects human poisoning or health emergencies to clinical helplines."
        ],
        badge="Chemical Safety"
    )

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Weather-Linked Risk Engine",
        [
            "Live Open-Meteo Integration: Checks wind speed, rain probability & temperature before recommending sprays.",
            "Drift Prevention: Warns farmers if high winds (>15 km/h) would cause spray drift.",
            "Wash-Off Warning: Alerts against fungicide/pesticide application if rain is forecast within 6-12 hours.",
            "Heat Degradation: Identifies midday heat risks for chemical volatility."
        ],
        badge="Environmental Safety"
    )

    # -------------------------------------------------------------
    # SLIDE 7: Web Application & Extension Officer Tools
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Web Application Features & Extension Dashboard")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "End-to-End Voice Chat",
        [
            "Seamless Mic Recording with AudioContext & MediaRecorder fallbacks.",
            "Instant transcript preview with user editing support.",
            "Spoken neural audio playback with auto-play & replay controls.",
            "Audio file uploader for testing voice without microphone hardware."
        ],
        badge="Farmer Chat"
    )

    add_card(s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Decision Tools",
        [
            "Yield & Fertilizer Calculator: Computes Urea & NPS requirements for target yields.",
            "Weather Risk Bar: Real-time spraying conditions for Adama, Hawassa, Bahir Dar, etc.",
            "Multi-language switcher: Amharic, Afaan Oromo, Tigrinya, Somali, English."
        ],
        badge="Agronomic Tools"
    )

    add_card(s7, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Extension Analytics",
        [
            "Live Hotline 8028 Call Monitor: Inspects call duration, caller region & sentiment.",
            "Pest Outbreak Tracking: Aggregates symptom queries across woredas.",
            "Automated Advisory SMS: Dispatches tailored SMS follow-ups to callers."
        ],
        badge="Management"
    )

    # -------------------------------------------------------------
    # SLIDE 8: Privacy, Telephony & Production Readiness
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Privacy by Design & Telephony Engineering")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Farmer Privacy & Data Protection",
        [
            "Salted Caller Hashing: Raw phone numbers are NEVER stored in plain text.",
            "SHA-256 with rotating server salt generates pseudonymous caller IDs.",
            "Zero PII Leakage: Database records store anonymized session logs.",
            "Farmer Consent Gate: Caller is prompted for data storage permission on first turn.",
            "Encrypted database connections and Redis state caching."
        ],
        badge="Security & Compliance"
    )

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Production Telephony Infrastructure",
        [
            "Asterisk PBX Container with custom AudioSocket application gateway.",
            "Multi-Turn Session Manager: Maintains conversational context across turns.",
            "Sub-2s End-to-End Latency: Parallel audio streaming and LLM token pipeline.",
            "Docker Compose Deployment: 6 container microservice architecture (Frontend, Backend, Worker, Redis, Postgres, Asterisk)."
        ],
        badge="Scalability"
    )

    # -------------------------------------------------------------
    # SLIDE 9: Project Impact & Future Roadmap
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Impact, Current Milestones & Future Vision")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Accomplished Milestones",
        [
            "End-to-end voice chat running in browser and Asterisk 8028 hotline.",
            "Full RAG knowledge base grounded in Ethiopia MoA manuals.",
            "8/8 automated integration test suite verified in Docker.",
            "5-language translation system implemented."
        ],
        badge="Current Status"
    )

    add_card(s9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Fine-Tuned Dialect Models",
        [
            "Active model fine-tuning underway for regional Ethiopian dialects.",
            "Custom acoustic models in amharic-stt-2 repository.",
            "Adapting to regional accents: Wollo, Gojjam, Shewa, and Harar.",
            "Reduces Word Error Rate (WER) in noisy agricultural environments."
        ],
        badge="Under Development"
    )

    add_card(s9, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Future Roadmap",
        [
            "Vision AI: Camera-based crop pest & leaf disease diagnosis in web app.",
            "Solar-Powered Edge Box: Deploying Ollama & Whisper in off-grid woredas.",
            "Integration with National Agricultural Research Institutes (EIAR).",
            "Expansion to additional languages: Sidama, Wolaytta, Afar."
        ],
        badge="Next Steps"
    )

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Q&A
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = C_PRIMARY_DARK
    bg10.line.fill.background()

    s10_box = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf10 = s10_box.text_frame
    tf10.word_wrap = True

    p_end_badge = tf10.paragraphs[0]
    p_end_badge.text = "HELLO FARMER • ሄሎ ፋርመር"
    p_end_badge.font.size = Pt(14)
    p_end_badge.font.bold = True
    p_end_badge.font.color.rgb = C_ACCENT_GOLD
    p_end_badge.space_after = Pt(12)

    p_end_h1 = tf10.add_paragraph()
    p_end_h1.text = "Transforming Agriculture One Voice at a Time"
    p_end_h1.font.size = Pt(40)
    p_end_h1.font.bold = True
    p_end_h1.font.color.rgb = C_WHITE
    p_end_h1.space_after = Pt(16)

    p_end_sub = tf10.add_paragraph()
    p_end_sub.text = "Thank you! Questions, Feedback & Demonstrations Welcome.\n\nProject Team: Yeabsira Girma • KalMesKeb • Yohannes Girma"
    p_end_sub.font.size = Pt(18)
    p_end_sub.font.color.rgb = RGBColor(209, 231, 221)

    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Hello_Farmer_Presentation.pptx"
    create_deck(out)
